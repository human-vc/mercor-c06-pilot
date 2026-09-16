import os

import subprocess

import sys

from pathlib import Path


MAX_CHARS = 20000

MAX_LIST = 400

SKIP_SUFFIXES = {".mp4", ".mov", ".png", ".jpg", ".jpeg", ".gif", ".zip"}

PY_TIMEOUT = 180


def declarations():

    return [

        {"name": "list_files", "description": "List the files available for this task with sizes in bytes. Optionally restrict to a folder prefix.", "parameters": {"type": "object", "properties": {"path": {"type": "string", "description": "folder prefix, empty for everything"}}}},

        {"name": "read_file", "description": "Read a task file as text (txt, md, csv, xlsx, docx, pptx, pdf). Long files are returned in chunks of 20000 characters; pass offset to read the next chunk. For large spreadsheets prefer run_python with openpyxl.", "parameters": {"type": "object", "properties": {"path": {"type": "string"}, "offset": {"type": "integer", "description": "character offset to start from, default 0"}}, "required": ["path"]}},

        {"name": "run_python", "description": "Run a Python 3 script in the task folder and return its stdout and stderr (openpyxl, pandas, python-docx, python-pptx, pypdf available). Each call starts a fresh interpreter: variables and imports do not persist between calls, so make every script self-contained. Use it for calculations and for reading spreadsheets cell by cell. Files it writes into the task folder count as deliverables.", "parameters": {"type": "object", "properties": {"code": {"type": "string"}}, "required": ["code"]}},

        {"name": "write_file", "description": "Write or overwrite a deliverable file with text content.", "parameters": {"type": "object", "properties": {"path": {"type": "string"}, "content": {"type": "string"}}, "required": ["path", "content"]}},

        {"name": "message_client", "description": "Send one message to the client to ask about a requirement. You may use this at most once; the client's reply is final.", "parameters": {"type": "object", "properties": {"message": {"type": "string"}}, "required": ["message"]}},

    ]


def read_as_text(path):

    p = Path(path)

    suf = p.suffix.lower()

    if suf in (".txt", ".md", ".csv", ".json", ".html"):

        return p.read_text(errors="replace")

    if suf == ".xlsx":

        import openpyxl

        wb = openpyxl.load_workbook(p, data_only=True, read_only=True)

        out = []

        for ws in wb.worksheets:

            out.append(f"## sheet {ws.title}")

            for row in ws.iter_rows(values_only=True):

                if any(v is not None for v in row):

                    out.append("\t".join("" if v is None else str(v) for v in row))

        return "\n".join(out)

    if suf == ".docx":

        import docx

        d = docx.Document(str(p))

        parts = [para.text for para in d.paragraphs]

        for t in d.tables:

            for row in t.rows:

                parts.append("\t".join(c.text for c in row.cells))

        return "\n".join(parts)

    if suf == ".pptx":

        from pptx import Presentation

        parts = []

        for i, slide in enumerate(Presentation(str(p)).slides, 1):

            parts.append(f"## slide {i}")

            for shape in slide.shapes:

                if shape.has_text_frame:

                    parts.append(shape.text_frame.text)

        return "\n".join(parts)

    if suf == ".pdf":

        import logging

        logging.getLogger("pypdf").setLevel(logging.ERROR)

        import pypdf

        return "\n".join(pg.extract_text() or "" for pg in pypdf.PdfReader(str(p)).pages)

    return p.read_bytes()[:2000].decode("latin-1")


class Toolbox:

    def __init__(self, workdir):

        self.workdir = Path(workdir).resolve()

        self.writes = []

        self.baseline = self._snapshot()


    def _snapshot(self):

        return {str(p.relative_to(self.workdir)): (p.stat().st_size, p.stat().st_mtime_ns) for p in self.workdir.rglob("*") if p.is_file()}


    def changed_files(self):

        now = self._snapshot()

        changed = [k for k, v in now.items() if self.baseline.get(k) != v]

        return sorted(set(changed) | set(self.writes))


    def _safe(self, rel):

        p = (self.workdir / rel).resolve()

        if self.workdir not in p.parents and p != self.workdir:

            raise ValueError("path escapes workdir")

        return p


    def list_files(self, path=""):

        base = self._safe(path) if path else self.workdir

        if not base.exists():

            return f"ERROR: no such folder {path}"

        rows = []

        for p in sorted(base.rglob("*")):

            if p.is_file() and not p.name.startswith(".") and p.suffix.lower() not in SKIP_SUFFIXES:

                rows.append(f"{p.relative_to(self.workdir)}\t{p.stat().st_size}")

        if len(rows) > MAX_LIST:

            rows = rows[:MAX_LIST] + [f"[{len(rows) - MAX_LIST} more files; pass a folder prefix to narrow]"]

        return "\n".join(rows) if rows else "(no files)"


    def read_file(self, path, offset=0):

        p = self._safe(path)

        if not p.exists():

            return f"ERROR: no such file {path}"

        try:

            text = read_as_text(p)

        except Exception as e:

            return f"ERROR: could not read {path}: {e}"

        offset = int(offset or 0)

        chunk = text[offset:offset + MAX_CHARS]

        if offset + MAX_CHARS < len(text):

            chunk += f"\n[TRUNCATED at {offset + MAX_CHARS} of {len(text)} chars; call again with offset={offset + MAX_CHARS}]"

        return chunk


    def run_python(self, code):

        env = {"PATH": "/usr/bin:/bin", "HOME": str(self.workdir), "PYTHONIOENCODING": "utf-8", "TMPDIR": str(self.workdir)}

        sandbox = str(Path(__file__).resolve().parent / "sandbox.py")

        try:

            r = subprocess.run([sys.executable, "-I", sandbox, str(self.workdir)], input=code, cwd=self.workdir, env=env, capture_output=True, text=True, timeout=PY_TIMEOUT)

        except subprocess.TimeoutExpired:

            return f"ERROR: script exceeded {PY_TIMEOUT} seconds"

        out = r.stdout

        if r.stderr:

            err = "\n".join(line for line in r.stderr.splitlines() if "sandbox.py" not in line)

            err = err.replace(str(Path(__file__).resolve().parent.parent), "<sandbox>").replace(sys.prefix, "<python>")

            out += ("\n" if out else "") + "STDERR:\n" + err

        if r.returncode != 0:

            out += f"\n[exit code {r.returncode}]"

        if len(out) > 12000:

            out = out[:6000] + "\n[... TRUNCATED ...]\n" + out[-6000:]

        return out or "(no output)"


    def write_file(self, path, content):

        p = self._safe(path)

        p.parent.mkdir(parents=True, exist_ok=True)

        p.write_text(content)

        self.writes.append(str(p.relative_to(self.workdir)))

        return f"wrote {path} ({len(content)} chars)"
